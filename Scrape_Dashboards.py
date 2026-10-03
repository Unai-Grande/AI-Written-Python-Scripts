#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NOC Dashboard Scraper - NagVis + Nagios Edition (multi-server)

Scrapes the configured NagVis/Nagios servers into one combined report.
Every row carries a "Nagios Server" column so downstream tools
(normalized_key_matrix.py, host_comparison_report.py) can filter by
server. On instance-style servers (see name2 below) the value embeds the
instance (e.g. "192.0.2.11/amicus") since each dashboard comes from a
separate per-instance Nagios install; single-instance servers keep the
plain server IP. name1 uses a verified map->hostgroup table
(DASHBOARDS_NAME1); name3 auto-discovers its hostgroups live until a
verified table is added - see the SERVERS comment.

Requirements:
    pip install requests beautifulsoup4 openpyxl lxml
"""

import re
import time
import socket
import threading
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

#  Config 
# All active NagVis / Nagios servers to scrape, IP -> settings.
# (IPs below are RFC 5737 documentation addresses - replace with your own.)
#
# "dashboards" can be:
#   - the name of a verified {label: {"map":..., "groups":[...]}} table
#     defined in this module (like DASHBOARDS_NAME1 below, verified from
#     the server's objectjson.cgi);
#   - "auto": discover the server's hostgroups live via objectjson.cgi and
#     treat each hostgroup as its own dashboard entry. Auto mode guarantees
#     full host coverage for the comparison workflow, but the Dashboard
#     labels are raw hostgroup names rather than curated NagVis map names -
#     replace "auto" with a verified table once the map -> hostgroup
#     relationships have been confirmed against objectjson.cgi;
#   - "instances": for servers that run a SEPARATE Nagios instance per
#     dashboard under /usr/local/nagios/<name>. In the typical Apache
#     layout, each instance's CGIs are ScriptAliased at /nagios_<inst>
#     mapping DIRECTLY to sbin - so status.cgi lives at
#     http://<ip>/nagios_<inst>/status.cgi with NO /cgi-bin/ segment - and
#     its web UI is Aliased at /<inst>.
#     The "instances" list makes each instance its own dashboard entry; the
#     instance's CGI directory is probed from "cgi_path_templates" ({inst}
#     replaced with the instance name, status.cgi requested directly under
#     it, first template answering 200 wins - the /nagios_{inst} shape is
#     listed first).
SERVERS = {
    "192.0.2.10": {"name": "name1", "dashboards": "DASHBOARDS_NAME1"},
    "192.0.2.11": {"name": "name2", "dashboards": "instances",
                   "instances": ["amicus", "ciao", "cogito", "ergo", "sum"],
                   "cgi_path_templates": ["/nagios_{inst}", "/nagios/{inst}/cgi-bin",
                                          "/{inst}/nagios/cgi-bin", "/{inst}/cgi-bin"]},
    "192.0.2.12": {"name": "name3", "dashboards": "auto"},
}
SERVER_IPS = set(SERVERS.keys())

# Pattern matching your hostname naming convention. Used to skip hostname
# cells when picking the service-name cell on status pages. Adjust to fit
# your environment (e.g. r'^(SRV|srv)\d' for hosts named SRV01..., SRV02...).
HOSTNAME_PATTERN = re.compile(r'^(SRV|srv)\d|^\d{1,3}\.\d')


def base_url(server_ip):
    return f"http://{server_ip}"


def default_cgi_base(server_ip):
    """The standard single-instance layout: http://<ip>/nagios/cgi-bin"""
    return f"{base_url(server_ip)}/nagios/cgi-bin"


def nagios_hg_url(cgi_base, hg):
    if hg == "__ALL__":
        # Instance-style entries have no hostgroup mapping - pull every host.
        return (f"{cgi_base}/status.cgi"
                f"?host=all&style=hostdetail&limit=0&sorttype=1&sortoption=1")
    return (f"{cgi_base}/status.cgi"
            f"?hostgroup={requests.utils.quote(hg, safe='')}"
            f"&style=hostdetail&limit=0&sorttype=1&sortoption=1")


def nagios_host_url(cgi_base, host):
    return f"{cgi_base}/status.cgi?host={host}&style=detail&limit=0"


def nagios_extinfo_url(cgi_base, host):
    return f"{cgi_base}/extinfo.cgi?type=1&host={host}"


def probe_instance_url(server_ip, inst, templates):
    """Try each CGI path template until the instance's status.cgi answers
    with 200. Templates name the CGI directory itself (status.cgi requested
    DIRECTLY under it - matching the /nagios_<inst> ScriptAlias layout).
    Returns the full CGI base (e.g. 'http://<ip>/nagios_<inst>') or None."""
    for tpl in templates:
        path = tpl.format(inst=inst)
        url = f"http://{server_ip}{path}/status.cgi"
        try:
            r = get_session().get(url, params={"hostgroup": "all", "style": "overview"},
                                  timeout=8, verify=False)
            if r.status_code == 200:
                return f"http://{server_ip}{path}"
            if r.status_code in (401, 403):
                print(f"    [WARN] {server_ip}/{inst}: {path} exists but needs auth "
                      f"(HTTP {r.status_code}) - instance skipped here; "
                      f"nagios_full_export.py's SSH path still covers it.")
        except Exception:
            continue
    return None


MAP_WORKERS  = 4
HOST_WORKERS = 8

IP_RE        = re.compile(r'\b(\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})\b')
DURATION_RE  = re.compile(r'\d+d\s+\d+h|\d+h\s+\d+m|\d+m\s+\d+s')
VALID_STATUS = {"UP", "DOWN", "UNREACHABLE", "PENDING", "UNKNOWN", "WARNING",
                "OK", "CRITICAL"}
STATE_RANK   = {"UP": 0, "OK": 0, "PENDING": 1, "UNKNOWN": 2,
                "WARNING": 3, "DOWN": 4, "CRITICAL": 4, "UNREACHABLE": 4}

#  Example verified hostgroup mapping for name1 
# Replace with your own map -> hostgroup relationships, confirmed against
# the server's objectjson.cgi. "map" is the NagVis map name; "groups" are
# the Nagios hostgroups whose hosts belong on that dashboard.
DASHBOARDS_NAME1 = {
    "Web Frontend Dashboard": {
        "map":    "WEB_FRONTEND",
        "groups": ["WEB-1", "WEB-2", "WEB-3", "WEB-STANDBY"],
    },
    "Database Dashboard": {
        "map":    "DATABASE",
        "groups": ["SQL-PRIMARY", "SQL-REPLICA"],
    },
    "Middleware Dashboard": {
        "map":    "MIDDLEWARE",
        "groups": ["MQ Services", "API Gateway"],
    },
    "Authentication Dashboard": {
        "map":    "AUTH",
        "groups": ["LDAP", "SSO"],
    },
    "Core Infrastructure Dashboard": {
        "map":    "CORE_INFRA",
        "groups": ["Core Monitoring", "NTP"],
    },
    "Staging Dashboard": {
        "map":    "STAGING",
        "groups": ["WEB-STAGING", "TEST"],
    },
}

#  Thread-local HTTP session 
_local = threading.local()

def get_session():
    if not hasattr(_local, "session"):
        s = requests.Session()
        s.headers.update({
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        })
        _local.session = s
    return _local.session


#  Status helpers 
def norm(raw):
    if not raw:
        return ""
    u = str(raw).strip().upper()
    if u in ("OK", "UP"):                        return "UP"
    if u in ("DOWN", "CRITICAL", "UNREACHABLE"): return "DOWN"
    if u in ("WARNING", "WARN"):                 return "WARNING"
    if u == "UNKNOWN":                           return "UNKNOWN"
    if u == "PENDING":                           return "PENDING"
    return u

def worst(states):
    ranked = sorted([s for s in states if s],
                    key=lambda s: STATE_RANK.get(s, 0), reverse=True)
    return ranked[0] if ranked else ""


#  Fetch hostgroup page 
def fetch_hostgroup(hg_name, cgi_base):
    url = nagios_hg_url(cgi_base, hg_name)
    try:
        r = get_session().get(url, timeout=30, verify=False)
        r.raise_for_status()
    except Exception as e:
        print(f"    [WARN] hostgroup '{hg_name}': {e}")
        return {}

    soup    = BeautifulSoup(r.text, "lxml")
    results = {}

    for table in soup.find_all("table", class_="status"):
        for row in table.find_all("tr"):
            tds      = row.find_all("td")
            td_texts = [td.get_text(" ", strip=True) for td in tds]
            if len(td_texts) < 3:
                continue

            # Only rows with an exact status word
            status_idx = None
            for i, txt in enumerate(td_texts):
                if txt.strip().upper() in VALID_STATUS:
                    status_idx = i
                    break
            if status_idx is None:
                continue

            # Hostname from extinfo anchor
            hostname = ""
            for td in tds:
                a = td.find("a", href=re.compile(r"extinfo\.cgi\?type=1&host="))
                if a:
                    m = re.search(r"host=([^&\"]+)", a["href"])
                    if m:
                        hostname = m.group(1).strip()
                    break
            if not hostname or hostname in results:
                continue

            host_status = norm(td_texts[status_idx])
            duration    = ""
            status_info = ""

            if status_idx + 2 < len(td_texts):
                cand = td_texts[status_idx + 2]
                if DURATION_RE.search(cand):
                    duration = cand
            if status_idx + 3 < len(td_texts):
                status_info = td_texts[status_idx + 3]

            results[hostname] = {
                "host_status": host_status,
                "duration":    duration,
                "status_info": status_info,
            }

    return results


def get_hosts_for_dashboard(groups, cgi_base):
    merged = {}
    for hg in groups:
        for hostname, info in fetch_hostgroup(hg, cgi_base).items():
            if hostname not in merged:
                merged[hostname] = info
    return merged


#  Service detail 
def get_host_service_detail(hostname, cgi_base):
    """
    Actual column layout (verified from live debug output):

    First service row (16 cols):
      [host x3][empty x3][svc_name x3][empty x2][STATUS][lastcheck][duration][attempt][info]
    Subsequent rows (11 cols):
      [empty][svc_name x3][empty x2][STATUS][lastcheck][duration][attempt][info]

    Strategy: find the first cell that is exactly a known status word.
    Service name = last non-empty, non-hostname cell before the status cell.
    Duration = status+2, info = status+4.
    """
    url = nagios_host_url(cgi_base, hostname)
    try:
        r = get_session().get(url, timeout=20, verify=False)
        r.raise_for_status()
    except Exception as e:
        print(f"    [WARN] service detail '{hostname}': {e}")
        return [], ""

    soup     = BeautifulSoup(r.text, "lxml")
    services = []

    for table in soup.find_all("table"):
        for row in table.find_all("tr"):
            tds      = row.find_all("td")
            td_texts = [td.get_text(" ", strip=True) for td in tds]
            if len(td_texts) < 8:
                continue

            # Find the status cell index
            status_idx = None
            for i, txt in enumerate(td_texts):
                if txt.strip().upper() in VALID_STATUS:
                    status_idx = i
                    break
            if status_idx is None:
                continue

            # Service name: last non-empty, non-hostname cell before status
            svc_name = ""
            for j in range(status_idx - 1, -1, -1):
                candidate = td_texts[j].strip()
                if candidate and candidate.lower() not in ("service", "host", ""):
                    if not HOSTNAME_PATTERN.match(candidate):
                        svc_name = candidate
                        break
            if not svc_name:
                continue

            svc_state = norm(td_texts[status_idx])
            svc_dur   = td_texts[status_idx + 2] if status_idx + 2 < len(td_texts) else ""
            svc_info  = td_texts[status_idx + 4] if status_idx + 4 < len(td_texts) else (
                        td_texts[status_idx + 3] if status_idx + 3 < len(td_texts) else "")

            if svc_dur and not DURATION_RE.search(svc_dur):
                svc_dur = ""

            services.append({
                "name":  svc_name,
                "state": svc_state,
                "dur":   svc_dur,
                "info":  svc_info.strip()[:200],
            })

    # Deduplicate services by name (keep first occurrence)
    seen_svcs = set()
    unique_services = []
    for s in services:
        if s["name"] not in seen_svcs:
            seen_svcs.add(s["name"])
            unique_services.append(s)

    # IP: prefer internal RFC-1918, skip the server itself
    ip = ""
    for m in IP_RE.finditer(soup.get_text(" ")):
        c = m.group(1)
        if c.startswith(("10.", "172.", "192.168.")) and c not in SERVER_IPS:
            ip = c
            break

    # If no IP found from page text, try DNS resolution.
    # Works for both short hostnames and FQDNs on the internal network.
    if not ip:
        try:
            ip = socket.gethostbyname(hostname)
        except socket.gaierror:
            pass

    return unique_services, ip




#  Fetch canonical IP and full service list from Nagios extinfo 
def get_nagios_extinfo(hostname, cgi_base):
    """
    Fetch extinfo.cgi?type=1&host=X to get the canonical host IP address
    as Nagios has it configured (the 'Address' field), which is more reliable
    than scanning the service detail page text.

    If the Address field contains an FQDN rather than a bare IP, resolve it
    via DNS. Falls back to DNS resolution of the hostname itself if extinfo
    yields nothing useful.
    """
    nagios_ip = ""
    try:
        r = get_session().get(
            nagios_extinfo_url(cgi_base, hostname),
            timeout=15, verify=False)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "lxml")

        # Nagios extinfo: label td (dataVar) | value td (dataVal)
        for td in soup.find_all("td", class_="dataVar"):
            if "address" in td.get_text(strip=True).lower():
                val_td = td.find_next_sibling("td", class_="dataVal")
                if val_td:
                    candidate = val_td.get_text(strip=True)
                    # Already an IP address - use directly
                    if IP_RE.match(candidate):
                        nagios_ip = candidate
                        break
                    # FQDN in the Address field - resolve it
                    if candidate and "." in candidate:
                        try:
                            nagios_ip = socket.gethostbyname(candidate)
                        except socket.gaierror:
                            pass
                        break

        # Fallback 1: scan full page for first internal RFC-1918 IP
        if not nagios_ip:
            for m in IP_RE.finditer(soup.get_text(" ")):
                c = m.group(1)
                if c.startswith(("10.", "172.", "192.168.")) and c not in SERVER_IPS:
                    nagios_ip = c
                    break

    except Exception:
        pass

    # Fallback 2: DNS-resolve the hostname itself if still no IP found.
    # Applies to both short hostnames (e.g. SRV01DB01) and FQDNs.
    # Short hostnames resolve via the internal DNS search domain.
    if not nagios_ip:
        try:
            nagios_ip = socket.gethostbyname(hostname)
        except socket.gaierror:
            pass

    return nagios_ip

#  Process one host 
def process_host(hostname, entry, map_url, cgi_base, server_label):
    host_status = entry["host_status"]
    host_dur    = entry["duration"]
    host_info   = entry["status_info"]

    services, ip = get_host_service_detail(hostname, cgi_base)
    nagios_ip     = get_nagios_extinfo(hostname, cgi_base)

    # Build full service list string: "ServiceName: STATE, ServiceName: STATE, ..."
    if services:
        all_svcs_str = ", ".join(
            f"{s['name']}: {s['state']}" for s in services
        )
    else:
        all_svcs_str = ""

    # problems is always defined here so the return dict can safely reference it
    problems      = []
    rolled_status = host_status
    duration      = host_dur
    status_info   = host_info

    if services:
        all_states    = [host_status] + [s["state"] for s in services]
        rolled_status = worst(all_states)

        problems = sorted(
            [s for s in services if s["state"] not in ("UP", "")],
            key=lambda s: STATE_RANK.get(s["state"], 0), reverse=True)

        if problems:
            worst_svc   = problems[0]
            duration    = worst_svc["dur"] or host_dur
            info_parts  = [
                f"{s['name']}: {s['state']}" +
                (f" - {s['info'][:80]}" if s["info"] else "")
                for s in problems[:5]
            ]
            status_info = " | ".join(info_parts)

    # Build problem service name: comma-joined names of all problem services
    problem_service = ", ".join(s["name"] for s in problems[:3]) if problems else ""

    return {
        "Dashboard":              "",
        "Dashboard URL":          map_url,
        "Nagios Server":          server_label,
        "Host":                   hostname,
        "Status":                 rolled_status,
        "Duration":               duration,
        "IP Address":             ip,
        "Found IP Address Nagios": nagios_ip,
        "Found Services Nagios":  all_svcs_str,
        "Problem Service":        problem_service,
        "Status Info":            status_info[:300],
        "Host Link":              nagios_host_url(cgi_base, hostname),
    }


#  Scrape one dashboard 
def scrape_dashboard(label, cfg, base, server_ip):
    map_name = cfg["map"]
    groups   = cfg["groups"]
    cgi_base = cfg.get("cgi_base") or default_cgi_base(server_ip)
    # "ip/instance" for per-instance servers (e.g. "192.0.2.11/amicus");
    # plain server_ip everywhere else (no instance concept).
    inst = cfg.get("instance", "")
    srv_label = f"{server_ip}/{inst}" if inst else server_ip
    if map_name:
        map_url = (f"{base}/nagvis/frontend/nagvis-js/index.php"
                   f"?mod=Map&act=view&show={map_name}&header_menu=0")
    else:
        # Auto-discovered hostgroup / per-dashboard instance with no verified
        # NagVis map: point the Dashboard URL at the instance's own web UI
        # (Alias /<inst> on per-instance servers), falling back to the
        # server's NagVis frontend. Either way it embeds the server IP, so
        # Nagios Server resolution keeps working.
        map_url = cfg.get("ui_url") or f"{base}/nagvis/frontend/nagvis-js/index.php"

    host_map = get_hosts_for_dashboard(groups, cgi_base)
    print(f"    '{label}': {len(host_map)} host(s) across {len(groups)} group(s)")

    if not host_map:
        return []

    records = []
    with ThreadPoolExecutor(max_workers=HOST_WORKERS) as ex:
        futures = {
            ex.submit(process_host, hn, info, map_url, cgi_base, srv_label): hn
            for hn, info in host_map.items()
        }
        for fut in as_completed(futures):
            try:
                rec = fut.result()
                rec["Dashboard"] = label
                records.append(rec)
            except Exception as e:
                print(f"    [WARN] host error: {e}")

    records.sort(key=lambda r: r.get("Host", "").lower())
    return records


#  Auto-discovery for servers without a verified dashboard mapping 
def discover_dashboards(base, server_label):
    """
    Build a {label: {"map": None, "groups": [hg]}} dict for a server whose
    NagVis map -> hostgroup relationships haven't been verified yet, by
    pulling its live hostgroup list. Each hostgroup becomes its own
    dashboard entry (label suffixed '(auto)'), which guarantees full host
    coverage even though the grouping isn't the curated map layout.

    Primary source: objectjson.cgi?query=hostgrouplist (clean JSON).
    Fallback: scrape hostgroup names out of status.cgi overview links.
    """
    discovered = {}
    try:
        r = get_session().get(
            f"{base}/nagios/cgi-bin/objectjson.cgi",
            params={"query": "hostgrouplist"},
            timeout=20, verify=False)
        r.raise_for_status()
        data = r.json().get("data", {}).get("hostgrouplist", None)
        if data:
            names = list(data.keys()) if isinstance(data, dict) else list(data)
            for hg in sorted(names):
                discovered[f"{hg} (auto)"] = {"map": None, "groups": [hg]}
    except Exception as e:
        print(f"    [WARN] {server_label}: hostgroup JSON discovery failed ({e}); "
              f"trying status.cgi overview fallback.")

    if not discovered:
        try:
            r = get_session().get(
                f"{base}/nagios/cgi-bin/status.cgi",
                params={"hostgroup": "all", "style": "overview"},
                timeout=30, verify=False)
            r.raise_for_status()
            soup = BeautifulSoup(r.text, "lxml")
            names = set()
            for a in soup.find_all("a", href=True):
                m = re.search(r"hostgroup=([^&\"]+)", a["href"])
                if m and m.group(1) != "all":
                    names.add(requests.utils.unquote(m.group(1)))
            for hg in sorted(names):
                discovered[f"{hg} (auto)"] = {"map": None, "groups": [hg]}
        except Exception as e:
            print(f"    [WARN] {server_label}: overview fallback failed too ({e}).")

    print(f"    {server_label}: auto-discovered {len(discovered)} hostgroup(s)")
    return discovered


def dashboards_for_server(server_ip, cfg, base):
    """Resolve a server's 'dashboards' setting to a concrete dict."""
    setting = cfg["dashboards"]
    if setting == "instances":
        # Per-dashboard Nagios instances: each instance becomes one
        # dashboard entry pulling ALL of its hosts, via whichever probed
        # URL path answers.
        entries = {}
        templates = cfg.get("cgi_path_templates", ["/nagios_{inst}"])
        for inst in cfg.get("instances", []):
            cgi_base = probe_instance_url(server_ip, inst, templates)
            if not cgi_base:
                print(f"    [WARN] {cfg['name']}/{inst}: no working CGI path found - "
                      f"instance skipped here (nagios_full_export.py's SSH path "
                      f"still covers it).")
                continue
            entries[inst] = {"map": None, "groups": ["__ALL__"],
                             "cgi_base": cgi_base,
                             "ui_url": f"http://{server_ip}/{inst}/",
                             "instance": inst}
        print(f"    {cfg['name']}: {len(entries)}/{len(cfg.get('instances', []))} "
              f"instance(s) reachable over HTTP")
        return entries
    if setting == "auto":
        return discover_dashboards(base, cfg["name"])
    if isinstance(setting, str):
        resolved = globals().get(setting)
        if resolved is None:
            print(f"    [WARN] Unknown dashboards table '{setting}' for {server_ip}; "
                  f"falling back to auto-discovery.")
            return discover_dashboards(base, cfg["name"])
        return resolved
    return setting


#  Scrape all dashboards on all servers 
def scrape_all():
    all_recs = []

    for server_ip, cfg in SERVERS.items():
        base = base_url(server_ip)
        print("=" * 62)
        print(f"  {cfg['name']}  ({server_ip})")
        print("=" * 62)

        dashboards = dashboards_for_server(server_ip, cfg, base)
        if not dashboards:
            print(f"  [WARN] No dashboards resolvable for {server_ip} - server "
                  f"unreachable or empty. Its hosts will be MISSING from the "
                  f"combined report.\n")
            continue

        items     = list(dashboards.items())
        completed = 0
        print(f"Scraping {len(items)} dashboards  ({MAP_WORKERS} parallel)\n")

        with ThreadPoolExecutor(max_workers=MAP_WORKERS) as ex:
            futures = {
                ex.submit(scrape_dashboard, lbl, dcfg, base, server_ip): lbl
                for lbl, dcfg in items
            }
            for fut in as_completed(futures):
                lbl = futures[fut]
                completed += 1
                try:
                    recs     = fut.result()
                    problems = sum(1 for r in recs if r.get("Status") not in ("UP",""))
                    all_recs.extend(recs)
                    print(f"  [{completed:02d}/{len(items)}] {lbl:<46}"
                          f"{len(recs):>5} host(s)   {problems} problem(s)")
                except Exception as e:
                    print(f"  [{completed:02d}/{len(items)}] {lbl} ERROR: {e}")
        print()

    all_recs.sort(key=lambda r: (r.get("Nagios Server",""), r.get("Dashboard",""),
                                 r.get("Host","").lower()))
    return all_recs


#  Excel 
COLUMNS = [
    ("Dashboard",               30),
    ("Dashboard URL",           44),
    ("Nagios Server",           14),
    ("Host",                    32),
    ("Status",                  12),
    ("Duration",                24),
    ("IP Address",              16),
    ("Found IP Address Nagios", 18),
    ("Found Services Nagios",   52),
    ("Problem Service",         28),
    ("Status Info",             58),
    ("Host Link",               48),
]

STATUS_STYLE = {
    "UP":      ("006100", "C6EFCE"),
    "DOWN":    ("9C0006", "FFC7CE"),
    "WARNING": ("7D4800", "FFEB9C"),
    "UNKNOWN": ("3F3F3F", "D9D9D9"),
    "PENDING": ("1F497D", "BDD7EE"),
}

HEADER_FILL = PatternFill("solid", start_color="1F3864")
TITLE_FILL  = PatternFill("solid", start_color="0D1F3C")
ALT_FILL    = PatternFill("solid", start_color="EEF3F8")
NO_FILL     = PatternFill(fill_type=None)
THIN        = Side(style="thin", color="BFBFBF")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER      = Alignment(horizontal="center", vertical="center")
LEFT        = Alignment(horizontal="left",   vertical="center", wrap_text=False)
HDR_FONT    = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT   = Font(name="Arial", size=10)
LINK_FONT   = Font(name="Arial", size=10, color="0563C1", underline="single")


def write_sheet(ws, records, title_text):
    col_names      = [c[0] for c in COLUMNS]
    status_col_idx = col_names.index("Status") + 1

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    tc = ws.cell(row=1, column=1, value=title_text)
    tc.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    tc.fill = TITLE_FILL; tc.alignment = CENTER
    ws.row_dimensions[1].height = 30

    for ci, (col_name, col_w) in enumerate(COLUMNS, 1):
        c = ws.cell(row=2, column=ci, value=col_name)
        c.font = HDR_FONT; c.fill = HEADER_FILL
        c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = col_w
    ws.row_dimensions[2].height = 22

    for ri, rec in enumerate(records, start=3):
        status = rec.get("Status", "")
        row_bg = ALT_FILL if ri % 2 == 0 else NO_FILL

        for ci, col_name in enumerate(col_names, 1):
            val  = rec.get(col_name, "")
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = BORDER; cell.alignment = LEFT

            if ci == status_col_idx:
                cell.alignment = CENTER
                if status in STATUS_STYLE:
                    fg, bg = STATUS_STYLE[status]
                    cell.font = Font(name="Arial", size=10, bold=True, color=fg)
                    cell.fill = PatternFill("solid", start_color=bg)
                else:
                    cell.font = Font(name="Arial", size=10, bold=True)
                    cell.fill = row_bg
            elif col_name in ("Dashboard URL", "Host Link") and val:
                cell.hyperlink = val
                cell.font = LINK_FONT; cell.fill = row_bg
            else:
                cell.font = BODY_FONT; cell.fill = row_bg

        ws.row_dimensions[ri].height = 15

    ws.freeze_panes = "A3"
    last_col = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A2:{last_col}{max(2, len(records)+2)}"


def build_excel(records, output_path):
    wb  = openpyxl.Workbook()
    now = datetime.now().strftime("%B %d, %Y  %H:%M")

    ws1 = wb.active; ws1.title = "Host Status"
    write_sheet(ws1, records, f"NOC Dashboard Report  -  {now}")

    ws2 = wb.create_sheet("Summary")
    ws2.merge_cells("A1:K1")
    t = ws2.cell(row=1, column=1, value="NOC - Scrape Summary")
    t.font = Font(name="Arial", bold=True, size=14, color="FFFFFF")
    t.fill = TITLE_FILL; t.alignment = CENTER
    ws2.row_dimensions[1].height = 30

    def hdr(row, col, val):
        c = ws2.cell(row=row, column=col, value=val)
        c.font = HDR_FONT; c.fill = HEADER_FILL
        c.alignment = CENTER; c.border = BORDER

    for ci, h in enumerate(["Metric", "Value"], 1):
        hdr(2, ci, h)

    stats = [
        ("Servers scraped",    len(set(r.get("Nagios Server","") for r in records))),
        ("Dashboards scraped", len(set((r.get("Nagios Server",""), r.get("Dashboard","")) for r in records))),
        ("Total hosts",        len(records)),
        ("UP",                 sum(1 for r in records if r.get("Status") == "UP")),
        ("DOWN / Critical",    sum(1 for r in records if r.get("Status") == "DOWN")),
        ("WARNING",            sum(1 for r in records if r.get("Status") == "WARNING")),
        ("UNKNOWN",            sum(1 for r in records if r.get("Status") == "UNKNOWN")),
        ("Total problems",     sum(1 for r in records if r.get("Status") not in ("UP",""))),
        ("Generated",          datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for ri, (lbl, val) in enumerate(stats, start=3):
        c1 = ws2.cell(row=ri, column=1, value=lbl)
        c1.font = Font(name="Arial", size=11, bold=True)
        c1.border = BORDER; c1.alignment = LEFT
        c2 = ws2.cell(row=ri, column=2, value=val)
        c2.font = Font(name="Arial", size=11)
        c2.border = BORDER; c2.alignment = LEFT

    for ci, h in enumerate(["Dashboard","Total","UP","DOWN","WARN","Map URL"], 1):
        hdr(2, ci + 3, h)

    db_map = {}
    for rec in records:
        key = (rec.get("Nagios Server",""), rec["Dashboard"], rec["Dashboard URL"])
        db_map.setdefault(key, []).append(rec)

    for ri, ((srv, lbl, url), recs) in enumerate(sorted(db_map.items()), start=3):
        def dc(col, val, fill=None, link=None, ri=ri):
            c = ws2.cell(row=ri, column=col, value=val)
            c.font = LINK_FONT if link else Font(name="Arial", size=10)
            c.border = BORDER; c.alignment = LEFT
            if fill: c.fill = fill
            if link: c.hyperlink = link
        dc(4, f"{lbl}  [{srv}]" if srv else lbl)
        dc(5, len(recs))
        dc(6, sum(1 for r in recs if r.get("Status") == "UP"),
           fill=PatternFill("solid", start_color="C6EFCE"))
        dc(7, sum(1 for r in recs if r.get("Status") == "DOWN"),
           fill=PatternFill("solid", start_color="FFC7CE"))
        dc(8, sum(1 for r in recs if r.get("Status") == "WARNING"),
           fill=PatternFill("solid", start_color="FFEB9C"))
        dc(9, url, link=url)

    for col, w in zip(["A","B","C","D","E","F","G","H","I"],
                      [26,12,3,36,10,10,10,10,52]):
        ws2.column_dimensions[col].width = w

    wb.save(output_path)
    print(f"\n  Saved -> {output_path}")


#  CSV export 
def build_csv(records, output_path):
    import csv
    col_names = [c[0] for c in COLUMNS]
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=col_names, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
    print(f"  Saved -> {output_path}")


#  Entry point 
def main():
    date_str     = datetime.now().strftime("%Y-%m-%d")
    base_name    = f"application_dashboard_host_services_{date_str}"
    xlsx_path    = f"{base_name}.xlsx"
    csv_path     = f"{base_name}.csv"

    print("=" * 62)
    print("  NOC Dashboard Scraper  (service-aware, parallel, multi-server)")
    print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 62)

    t0      = time.time()
    records = scrape_all()
    elapsed = time.time() - t0

    print(f"\nTotal hosts   : {len(records)}")
    print(f"Elapsed       : {elapsed:.1f}s")
    print("Building Excel ...")
    build_excel(records, xlsx_path)
    print("Building CSV ...")
    build_csv(records, csv_path)
    print(f"Done!")
    print(f"  XLSX : {xlsx_path}")
    print(f"  CSV  : {csv_path}")
    print("=" * 62)


if __name__ == "__main__":
    main()